from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
CAPTURES = ROOT / "docs/assets/proposal-captures"
OUT = ROOT / "docs/FinPass_AI_공모전_시나리오_화면흐름.docx"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell, color="D9D9D9", size="6"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_font(run, name="Noto Sans CJK KR", size=10.5, bold=False, color="000000"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(16)
    set_font(p.add_run(text), size=9, color="526786")


def add_screenshot(doc, filename, caption, width):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(CAPTURES / filename), width=Inches(width))
    add_caption(doc, caption)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.7)
section.right_margin = Inches(0.7)

styles = doc.styles
styles["Normal"].font.name = "Noto Sans CJK KR"
styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Noto Sans CJK KR")
styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Noto Sans CJK KR")
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK KR")
styles["Normal"].font.size = Pt(10.5)
for style_name in ("Title", "Heading 1", "Heading 2"):
    style = styles[style_name]
    style.font.name = "Noto Sans CJK KR"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Noto Sans CJK KR")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Noto Sans CJK KR")
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK KR")
    style.font.color.rgb = RGBColor(0, 0, 0)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
set_font(title.add_run("FinPass AI 대출 오류 해결 시나리오"), name="Noto Sans CJK KR", size=25, bold=True)
subtitle = doc.add_paragraph()
set_font(subtitle.add_run("고객 화면과 상담 채널이 Journey Context로 연결되는 MVP 화면 흐름"), size=12, color="526786")
doc.add_paragraph(
    "이 문서는 개인사업자 대출 과정에서 본인 인증과 한도 조회 오류가 반복될 때 "
    "FinPass AI가 고객의 진행 맥락을 보존하고, 챗봇·콜센터·영업점 상담으로 "
    "자연스럽게 연결하는 경험을 공모전 심사 관점에서 설명한다."
)

doc.add_heading("핵심 시나리오", level=1)
doc.add_paragraph(
    "고객은 오류 코드를 다시 설명하지 않는다. 앱은 오류와 진행 단계를 기록하고, "
    "챗봇은 그 Context를 반영해 해결 방법과 상담 채널을 제안한다. 상담원은 "
    "Context Pass를 받아 고객이 멈춘 지점에서 업무를 이어간다."
)

flow = [
    ("01", "대출 신청", "개인사업자 대출 상품을 선택하고 신청을 시작"),
    ("02", "본인 인증 오류", "AUTH_TIMEOUT 발생, 재시도와 오류 이력 기록"),
    ("03", "한도 조회 오류", "NT004 발생, 서버 응답 지연 상황 안내"),
    ("04", "Journey 챗봇", "현재 단계·실패 근거·재시도 횟수를 반영해 답변"),
    ("05", "채널 연결", "콜센터 연결 또는 영업점 방문 안내 선택"),
    ("06", "상담 인계", "Context Pass로 상담원이 즉시 업무를 이어받음"),
    ("07", "업무 재개", "대체 소득증빙 등 해결책 적용 후 중단 단계부터 재개"),
]
table = doc.add_table(rows=0, cols=3)
table.autofit = False
widths = [0.55, 1.55, 4.95]
for number, name, detail in flow:
    cells = table.add_row().cells
    for index, width in enumerate(widths):
        cells[index].width = Inches(width)
        cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_border(cells[index])
    shade(cells[0], "2459D3")
    set_font(cells[0].paragraphs[0].add_run(number), size=10, bold=True, color="FFFFFF")
    cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(cells[1].paragraphs[0].add_run(name), size=10.5, bold=True, color="1D3769")
    set_font(cells[2].paragraphs[0].add_run(detail), size=10, color="526786")
doc.add_paragraph()

doc.add_heading("MVP 구현 범위", level=1)
scope = [
    ("고객 앱", "개인사업자 대출 신청, 진행 단계 저장, 본인 인증·한도 조회 오류 표시, 챗봇 질문"),
    ("Journey Context Layer", "이벤트·오류·재시도·현재 단계를 정규화하고 Context Pass 발급"),
    ("상담 채널", "콜센터와 영업점 상담원이 고객 Journey를 조회하고 해결책을 제시"),
    ("AI 기능", "Context Summary, Intent 추론, 실패 감지, 근거 기반 Next Best Action"),
    ("관리자", "단계별 실패율, 오류 코드, 상담 전환율, Journey 운영 Insight"),
]
table = doc.add_table(rows=1, cols=2)
for i, text in enumerate(("구성 요소", "실제 동작 범위")):
    cell = table.rows[0].cells[i]
    shade(cell, "1D3769")
    set_cell_border(cell)
    set_font(cell.paragraphs[0].add_run(text), size=10, bold=True, color="FFFFFF")
for left, right in scope:
    cells = table.add_row().cells
    for cell in cells:
        set_cell_border(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_font(cells[0].paragraphs[0].add_run(left), size=10, bold=True, color="1D3769")
    set_font(cells[1].paragraphs[0].add_run(right), size=10, color="526786")

doc.add_heading("서비스 구조도", level=1)
doc.add_paragraph("채널별 로그를 하나로 복제하지 않고, 현재 업무에 필요한 Context만 생성해 다음 채널로 전달한다.")
layers = [
    ("고객 채널", "모바일 웹 앱  |  Journey 진행  |  오류 알림  |  Context 챗봇  |  상담 채널 선택", "EAF1FF"),
    ("Financial Journey Context Layer", "Event Collector  →  Journey State  →  Failure Detection  →  Context Pass  →  Reverse Handoff", "DCE8FF"),
    ("상담 채널", "콜센터 Copilot  |  영업점 Copilot  |  고객 동의  |  상담 결과  |  업무 재개", "EAF8F0"),
    ("AI 및 데이터 기반", "PostgreSQL·pgvector  |  AI-Hub 상담·상품  |  BPI Process  |  Banking77  |  Synthetic Journey", "FFF3DF"),
]
architecture = doc.add_table(rows=0, cols=1)
for label, detail, fill in layers:
    cell = architecture.add_row().cells[0]
    shade(cell, fill)
    set_cell_border(cell)
    p = cell.paragraphs[0]
    set_font(p.add_run(label), size=11, bold=True, color="1D3769")
    p.add_run("\n")
    set_font(p.add_run(detail), size=10, color="526786")
    next_p = cell.add_paragraph()
    next_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if label != "AI 및 데이터 기반":
        set_font(next_p.add_run("↓"), size=15, bold=True, color="5675B8")

doc.add_heading("아키텍처 구성과 기능", level=1)
components = [
    ("Event API", "앱·챗봇·상담 채널의 행동을 append-only 이벤트로 저장", "Journey Timeline과 처리 근거"),
    ("Journey State", "이벤트 순서로 현재 단계와 상태를 계산", "중단 지점부터 재개"),
    ("Failure Detection", "오류 횟수와 재시도를 점수화", "상담이 필요한 실패를 구분"),
    ("Context Pass", "동의 범위 안의 최소 정보만 발급·만료 관리", "안전한 채널 전달"),
    ("AI Copilot", "Context Summary와 Intent를 생성하고 조치를 추천", "상담 설명 반복 감소"),
    ("Reverse Handoff", "상담 결과와 다음 단계를 Journey에 기록", "서류 제출 단계부터 재개"),
    ("Admin Analytics", "단계·오류·채널 전환을 집계", "서비스 병목 확인"),
]
component_table = doc.add_table(rows=1, cols=3)
for i, text in enumerate(("기능", "처리 방식", "사용자 가치")):
    cell = component_table.rows[0].cells[i]
    shade(cell, "1D3769")
    set_cell_border(cell)
    set_font(cell.paragraphs[0].add_run(text), size=10, bold=True, color="FFFFFF")
for row in components:
    cells = component_table.add_row().cells
    for cell in cells:
        set_cell_border(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for index, value in enumerate(row):
        set_font(cells[index].paragraphs[0].add_run(value), size=9.5, bold=index == 0, color="1D3769" if index == 0 else "526786")

doc.add_heading("사용자 이용 흐름", level=1)
doc.add_paragraph("고객 앱에서 시작한 금융업무가 오류 해결과 상담 이후 다시 고객 화면으로 돌아오는 전체 흐름이다.")
user_flow = [
    "고객 앱에서 개인사업자 대출 신청",
    "본인 인증 오류 AUTH_TIMEOUT 발생",
    "대출 한도 조회 오류 NT004 발생",
    "오류 이력과 현재 단계가 Journey에 저장",
    "챗봇 질문 및 Journey Context 기반 답변",
    "콜센터 연결 또는 영업점 방문 선택",
    "고객 동의 후 Context Pass 발급",
    "상담원이 실패 근거와 해결책 확인",
    "상담 결과 저장 및 서류 제출 단계부터 재개",
]
flow_table = doc.add_table(rows=0, cols=3)
for index, item in enumerate(user_flow, 1):
    cells = flow_table.add_row().cells
    for cell in cells:
        set_cell_border(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    shade(cells[0], "2459D3" if index % 2 else "DCE8FF")
    set_font(cells[0].paragraphs[0].add_run(f"{index:02d}"), size=10, bold=True, color="FFFFFF" if index % 2 else "1D3769")
    cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(cells[1].paragraphs[0].add_run(item), size=10.5, bold=True, color="1D3769")
    set_font(cells[2].paragraphs[0].add_run("다음 단계로 이동" if index < len(user_flow) else "Journey 재개 완료"), size=9.5, color="526786")

doc.add_heading("고객 앱 화면", level=1)
doc.add_paragraph("실제 은행 앱에서 볼 수 있는 오류 알림 구조를 참고해, 오류의 원인과 응답코드를 명확히 보여준다.")
add_screenshot(doc, "01-identity-error.png", "그림 1. 본인 인증 단계에서 AUTH_TIMEOUT 오류가 발생한 화면", 3.15)
add_screenshot(doc, "02-limit-error.png", "그림 2. 대출 한도 조회 단계에서 NT004 오류가 발생한 화면", 3.15)

doc.add_heading("Journey Context 기반 챗봇", level=1)
doc.add_paragraph(
    "챗봇은 일반적인 FAQ를 반복하지 않는다. 고객의 현재 단계가 본인 인증인지, "
    "한도 조회인지, 어떤 오류가 몇 회 발생했는지를 확인한 뒤 답변을 구성한다. "
    "문제가 반복되면 콜센터와 영업점 중 다음 채널을 선택할 수 있다."
)
add_screenshot(doc, "03-context-chat-handoff.png", "그림 3. 현재 Journey Context를 반영한 챗봇 답변과 채널 연결", 3.15)

doc.add_page_break()
doc.add_heading("상담 채널 인계 화면", level=1)
doc.add_paragraph(
    "고객이 채널을 바꾸더라도 상담원은 Context Pass를 통해 현재 업무 상태를 바로 확인한다. "
    "고객은 오류 코드와 진행 단계를 다시 설명하지 않고, 상담원은 대체 증빙이나 다음 업무를 안내한다."
)
add_screenshot(doc, "04-agent-copilot.png", "그림 4. 상담원의 Context Summary, 실패 근거, Next Best Action", 6.25)

doc.add_heading("운영 개선으로의 확장", level=1)
doc.add_paragraph(
    "관리자는 Journey 전체를 집계해 어느 단계에서 고객이 멈추는지 확인한다. "
    "개별 상담을 처리하는 Copilot과 서비스 병목을 찾는 Analytics가 동일한 Journey 이벤트를 사용한다."
)
add_screenshot(doc, "05-admin-analytics.png", "그림 5. 오류 구간과 상담 전환을 확인하는 관리자 화면", 6.25)

doc.add_heading("기획서 핵심 메시지", level=1)
messages = [
    "오류를 단순 실패 화면으로 끝내지 않고, 다음 행동으로 연결한다.",
    "챗봇 답변부터 상담원 화면까지 동일한 Journey Context를 공유한다.",
    "콜센터와 영업점 모두 고객이 멈춘 단계에서 업무를 이어받는다.",
    "상담 결과는 고객 앱에 반영되어 서류 제출 단계부터 재개된다.",
]
for message in messages:
    p = doc.add_paragraph(style="List Bullet")
    set_font(p.add_run(message), size=11)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
set_font(footer.add_run("FinPass AI · MVP Scenario"), size=8, color="7A879B")

doc.save(OUT)
print(OUT)
